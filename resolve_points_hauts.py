"""Turn a partner's raw list of high points (Excel) into a usable sites.csv.

A partner (SDIS, département, water utility…) typically hands over a
spreadsheet with one row per high point: a commune, the kind of structure
(water tower, pylon…) and sometimes a street address — but no coordinates and
no height. This script rebuilds both by cross-checking public sources:

  BAN (adresse.data.gouv.fr)   geocodes « address + commune » → a point
  BD TOPO IGN (geopf WFS)      `reservoir` / `pylone` structures with `hauteur`
  ANFR (SUP_SUPPORT.geojson)   radio supports: height, kind, operator
  Lidar HD MNH (geopf WMS)     measured height, as an independent arbiter

Geometry and height come from BD TOPO (the physical structure); ANFR is a
cross-check and covers pylons BD TOPO lacks. A water tower carrying no antenna
is absent from ANFR, hence the need for both.

The spreadsheet is read leniently: the header row is found by its labels
(« où », « quoi », « adresse », « gestionnaire »…), not by its position, and a
row's progress stage is read from its highlight colour (see COLORS).

Output: `sites.csv` in the format the rest of the chain expects, plus
`resolution.csv`, which records every decision (source, distance to the
geocoded point, rejected candidates) — read it by hand, column `alertes`,
before launching any computation.

Optional per-region inputs, under data/<region>/:

  commune_alias.json     {"OLD NAME": "MERGED NAME"} for communes nouvelles
  sites_installes.json   cameras already installed, whose position wins

    uv run python resolve_points_hauts.py --region example \
        --xlsx data/example/points_hauts.xlsx   # dept from scenarios.json
"""

from __future__ import annotations

from proj_env import use_bundled_proj

use_bundled_proj()

import argparse
import csv
import difflib
import json
import math
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

import numpy as np
import openpyxl
import rasterio
from pyproj import Transformer
from shapely.geometry import shape, Point

from analysis_shape.zones import region_dept

REPO = Path(__file__).resolve().parent

GEOPF_WFS = "https://data.geopf.fr/wfs/ows"
GEOPF_WMS = "https://data.geopf.fr/wms-r"
MNH_LAYER = "IGNF_LIDAR-HD_MNH_ELEVATION.ELEVATIONGRIDCOVERAGE.LAMB93"
BAN = "https://api-adresse.data.gouv.fr/search/"
GEO_API = "https://geo.api.gouv.fr"

# Row background colour → progress stage. The spreadsheet encodes each
# record's status with a highlight, not with a column.
COLORS = {
    "FF92D050": "vert",    # in service
    "FFFFC000": "orange",  # temporary
    "FFFFFF00": "jaune",   # agreement still to be signed
}

# Structure kinds, from SUP_NATURE.txt in the national ANFR export.
ANFR_CHATEAU_EAU = {"4"}
ANFR_PYLONE = {"11", "12", "21", "22", "23", "24", "25", "26", "33", "42", "43", "48"}

# Communes nouvelles: partner spreadsheets often keep the name a commune had
# before a merger, while the geo API only knows the merged name — and a short or
# misspelt name may match nothing. The table is specific to each partner list,
# so it lives in data/<region>/commune_alias.json, as {"OLD NAME": "MERGED
# NAME"} (keys and values compared after `norm()`), and is loaded by main().
COMMUNE_ALIAS: dict[str, str] = {}


def load_commune_alias(path: Path) -> dict[str, str]:
    """Optional alias table for `resolve_commune`; empty if the file is absent."""
    if not path.exists():
        return {}
    return {norm(k): norm(v)
            for k, v in json.loads(path.read_text(encoding="utf-8")).items()}

to_l93 = Transformer.from_crs("EPSG:4326", "EPSG:2154", always_xy=True).transform
to_wgs = Transformer.from_crs("EPSG:2154", "EPSG:4326", always_xy=True).transform


def norm(s) -> str:
    """Upper case, no accents, no punctuation — for comparing place names."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    return re.sub(r"[^A-Z0-9]+", " ", s).strip()


def get_json(url: str, params: dict | None = None, retries: int = 3):
    if params:
        url = f"{url}?{urllib.parse.urlencode(params)}"
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(url, timeout=120) as r:
                return json.load(r)
        except Exception as exc:  # noqa: BLE001
            if attempt == retries - 1:
                raise
            print(f"  ! {type(exc).__name__} — retrying", file=sys.stderr)
            time.sleep(2 * (attempt + 1))


# --------------------------------------------------------------------------
# reading the spreadsheet
# --------------------------------------------------------------------------

def read_xlsx(path: Path) -> list[dict]:
    """Headers are found by their label: title rows vary from one spreadsheet
    to the next, column names much less so."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]

    wanted = {
        "où": "commune", "quoi": "nature", "adresse": "adresse",
        "zones géographiques": "zone", "gestionnaire": "gestionnaire",
        "avancement": "avancement", "n°": "num", "coût": "cout", "cout": "cout",
    }
    header_row, cols = None, {}
    for r in range(1, min(ws.max_row, 30) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            key = str(ws.cell(r, c).value or "").strip().lower()
            if key in wanted:
                found[wanted[key]] = c
        if "commune" in found:
            header_row, cols = r, found
            break
    if header_row is None:
        raise SystemExit(f"{path.name}: no header row (column « où ») found")

    def fill(cell) -> str:
        f = cell.fill
        if f is None or f.patternType is None or f.fgColor.type != "rgb":
            return ""
        return f.fgColor.rgb or ""

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        commune = ws.cell(r, cols["commune"]).value
        if not commune or not str(commune).strip():
            continue
        rec = {k: ws.cell(r, c).value for k, c in cols.items()}
        # The highlight covers the whole row: the first coloured cell decides.
        etape = "blanc"
        for c in cols.values():
            if (tag := COLORS.get(fill(ws.cell(r, c)))):
                etape = tag
                break
        rec.update(ligne=r, etape=etape)
        rows.append(rec)
    return rows


# --------------------------------------------------------------------------
# reference datasets
# --------------------------------------------------------------------------

def load_communes(dept: str) -> tuple[dict, list]:
    communes = get_json(
        f"{GEO_API}/departements/{dept}/communes",
        {"fields": "nom,code,contour,centre", "format": "json"},
    )
    for c in communes:
        c["geom"] = shape(c["contour"])
    return {norm(c["nom"]): c for c in communes}, communes


def resolve_commune(name: str, index: dict):
    n = COMMUNE_ALIAS.get(norm(name), norm(name))
    if n in index:
        return index[n]
    # A merged commune's name often contains the old one; else, closest match.
    contains = [k for k in index if n in k]
    if len(contains) == 1:
        return index[contains[0]]
    close = difflib.get_close_matches(n, list(index), n=1, cutoff=0.85)
    return index[close[0]] if close else None


def wfs_features(typename: str, bbox_l93: tuple, page: int = 2000) -> list[dict]:
    """The geopf WFS caps a request at 5000 features and returns no reliable
    `numberMatched`, so page until a page comes back short."""
    out, start = [], 0
    while True:
        url = (
            f"{GEOPF_WFS}?SERVICE=WFS&VERSION=2.0.0&REQUEST=GetFeature"
            f"&TYPENAMES={typename}&OUTPUTFORMAT=application/json"
            f"&SRSNAME=EPSG:2154&COUNT={page}&STARTINDEX={start}"
            f"&BBOX={','.join(str(int(v)) for v in bbox_l93)},EPSG:2154"
        )
        feats = get_json(url).get("features", [])
        out.extend(feats)
        if len(feats) < page:
            return out
        start += page


def load_bdtopo(communes: list) -> dict:
    """BD TOPO structures within the département's extent, keyed by kind."""
    xs, ys = [], []
    for c in communes:
        minx, miny, maxx, maxy = c["geom"].bounds
        for lon, lat in ((minx, miny), (maxx, maxy)):
            x, y = to_l93(lon, lat)
            xs.append(x)
            ys.append(y)
    bbox = (min(xs) - 2000, min(ys) - 2000, max(xs) + 2000, max(ys) + 2000)

    result = {}
    for key, typename in (("reservoir", "BDTOPO_V3:reservoir"),
                          ("pylone", "BDTOPO_V3:pylone")):
        try:
            feats = wfs_features(typename, bbox)
        except Exception as exc:  # noqa: BLE001
            print(f"  ! BD TOPO {key} unavailable ({exc}) — carrying on without it", file=sys.stderr)
            feats = []
        items = []
        for f in feats:
            geom = shape(f["geometry"])  # already in L93
            pt = geom.centroid
            props = f["properties"]
            lon, lat = to_wgs(pt.x, pt.y)
            items.append({
                "x": pt.x, "y": pt.y, "lon": lon, "lat": lat,
                "nature": props.get("nature"),
                "hauteur": props.get("hauteur"),
                "z_sol": props.get("altitude_minimale_sol"),
                "z_toit": props.get("altitude_maximale_toit"),
                "cleabs": props.get("cleabs"),
                "etat": props.get("etat_de_l_objet"),
            })
        result[key] = items
        print(f"  BD TOPO {key:10s} : {len(items)} objects")
    return result


def load_anfr(path: Path, dept: str) -> list[dict]:
    """The etalab export is a geojson with one feature per line: stream it
    rather than load 77 MB to keep a single département."""
    out, seen = [], set()
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip().rstrip(",")
            if not line.startswith('{ "type": "Feature"'):
                continue
            f = json.loads(line)
            p = f["properties"]
            # The export drops the leading zero of INSEE codes: a 02xxx code
            # arrives as 2xxx, which `startswith("21")` assigns to the 21.
            # Pad both sides before comparing; Corsica (« 2A004 ») already has
            # five characters and is left untouched.
            code = str(p.get("COM_CD_INSEE") or "").strip()
            if not code or not code.zfill(5).startswith(str(dept).strip().zfill(2)):
                continue
            lon, lat = f["geometry"]["coordinates"]
            # Several stations share one support: keep a single entry.
            key = (round(lat, 5), round(lon, 5), p.get("SUP_NM_HAUT"))
            if key in seen:
                continue
            seen.add(key)
            x, y = to_l93(lon, lat)
            adr = " ".join(str(p[k]) for k in
                           ("ADR_LB_LIEU", "ADR_LB_ADD1", "ADR_LB_ADD2", "ADR_LB_ADD3")
                           if p.get(k))
            out.append({
                "x": x, "y": y, "lon": lon, "lat": lat,
                "hauteur": p.get("SUP_NM_HAUT"), "nat_id": str(p.get("NAT_ID")),
                "insee": str(p.get("COM_CD_INSEE")), "adresse": adr,
                "sup_id": p.get("SUP_ID"),
            })
    print(f"  ANFR dept {dept}   : {len(out)} distinct supports")
    return out


# --------------------------------------------------------------------------
# row-by-row resolution
# --------------------------------------------------------------------------

def hauteur_lidar(x: float, y: float, demi_cote: float = 20.0, res: float = 1.0):
    """Structure height measured by the lidar, in metres above ground.

    The MNH is already MNS − MNT: on a water tower its local maximum is the
    tower's height. It is an arbiter independent of ANFR and BD TOPO, which
    disagree by 10 to 20 m.

    The window is deliberately narrow (40 m) and takes the maximum rather than
    the centre value: the footprint's centroid may miss the shaft, and a
    neighbouring tree would enter a wider window.
    """
    size = max(3, int(2 * demi_cote / res))
    params = {
        "SERVICE": "WMS", "VERSION": "1.3.0", "REQUEST": "GetMap",
        "LAYERS": MNH_LAYER, "STYLES": "", "FORMAT": "image/geotiff",
        "CRS": "EPSG:2154",
        "BBOX": f"{x - demi_cote},{y - demi_cote},{x + demi_cote},{y + demi_cote}",
        "WIDTH": str(size), "HEIGHT": str(size),
    }
    url = f"{GEOPF_WMS}?{urllib.parse.urlencode(params)}"
    try:
        with urllib.request.urlopen(url, timeout=120) as r:
            payload = r.read()
        if payload[:4] not in (b"II*\x00", b"MM\x00*"):
            return None  # outside lidar coverage the WMS returns an error image
        with rasterio.io.MemoryFile(payload) as mem, mem.open() as src:
            window = src.read(1).astype(float)
    except Exception:  # noqa: BLE001 — a cross-check, never blocking
        return None
    valid = window[(window > 0) & (window < 200)]
    return round(float(np.max(valid)), 1) if valid.size else None


def geocode(adresse: str | None, commune: str, insee: str):
    """Return (x_l93, y_l93, label, score) or None. Without an address the
    geocoder falls back on the town hall, flagged by a zero score."""
    q = " ".join(filter(None, [str(adresse).strip() if adresse else "", commune]))
    q = re.sub(r'["\']', " ", q).strip()
    if not q:
        return None
    try:
        res = get_json(BAN, {"q": q, "citycode": insee, "limit": 1})
    except Exception:  # noqa: BLE001
        return None
    feats = res.get("features") or []
    if not feats:
        return None
    f = feats[0]
    lon, lat = f["geometry"]["coordinates"]
    x, y = to_l93(lon, lat)
    score = f["properties"].get("score", 0) if adresse else 0.0
    return x, y, f["properties"].get("label", ""), score


def dist(a_x, a_y, b) -> float:
    return math.hypot(a_x - b["x"], a_y - b["y"])


def adresse_ratio(a: str | None, b: str | None) -> float:
    a, b = norm(a), norm(b)
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a, b).ratio()


# Above this ratio the ANFR address and the spreadsheet's describe the same
# place. That is a safer signal than geocoding, which readily lands on a street
# with a similar name (« rue de l'Avenir » for « avenue de l'Océanie »).
SEUIL_ADRESSE = 0.62


def kind_of(nature: str | None) -> str:
    """What the spreadsheet asks for."""
    n = norm(nature)
    if "CHATEAU" in n or "RESERVOIR" in n:
        return "chateau_eau"
    if "PYLONE" in n or "MAT" in n or "TOUR" in n:
        return "pylone"
    return "inconnu"


def bdtopo_kind(nature: str | None) -> str:
    """What the BD TOPO object actually is. The nuance matters: a « réservoir
    d'eau ou château d'eau AU SOL » is 3 m tall and carries nothing — taking it
    instead of the neighbouring tower spoils the whole viewshed."""
    n = norm(nature)
    if n == "CHATEAU D EAU":
        return "chateau_eau"
    if "RESERVOIR" in n and "SOL" in n:
        return "reservoir_sol"
    if "RESERVOIR" in n:
        return "reservoir_industriel"
    return "autre"


def hauteur_of(o: dict):
    """`hauteur` is empty on ~7 % of structures; the altitudes recover it."""
    if o.get("hauteur"):
        return float(o["hauteur"]), "hauteur"
    toit, sol = o.get("z_toit"), o.get("z_sol")
    if toit is not None and sol is not None and toit > sol:
        return round(float(toit) - float(sol), 1), "altitudes"
    return None, None


# A BAN geocode without a street number often lands a few hundred metres from
# the structure; beyond this radius, matching no longer means anything.
RAYON_RECHERCHE_M = 3000
# Below this gap between the two best candidates the choice is a coin toss,
# so the row is flagged for review.
ECART_AMBIGU_M = 150


def choisir_hauteur(out: dict, preference: str) -> None:  # noqa: C901
    """Choose between the three heights available for one structure.

    They do not measure the same thing and differ by 12 to 18 m:

      ANFR      declared height of the support, antenna mast included
      BD TOPO   top of the tank, from photogrammetry
      MNH lidar top of whatever stands there, measured

    BD TOPO and lidar are independent and agree; ANFR stands alone and is
    always higher. For mounting a camera what counts is the structure's
    accessible top, not the tip of a radio mast — hence the « ouvrage »
    default. The same gap had been seen in another département between the
    conseil départemental's heights and ANFR; it points the same way.
    """
    # A height supplied by the operator overrides any inference.
    if out.get("hauteur_exploitation") is not None:
        out["hauteur_retenue"] = round(float(out["hauteur_exploitation"]), 1)
        out["hauteur_origine"] = "exploitation"
        return

    ordre = {
        "ouvrage": ("hauteur_bdtopo", "hauteur_lidar", "hauteur_anfr"),
        "anfr": ("hauteur_anfr", "hauteur_bdtopo", "hauteur_lidar"),
        "lidar": ("hauteur_lidar", "hauteur_bdtopo", "hauteur_anfr"),
    }[preference]
    libelle = {"hauteur_bdtopo": "BD TOPO", "hauteur_lidar": "MNH lidar",
               "hauteur_anfr": "ANFR"}
    for key in ordre:
        if out.get(key):
            out["hauteur_retenue"] = round(float(out[key]), 1)
            out["hauteur_origine"] = libelle[key]
            if key == "hauteur_bdtopo" and out.get("hauteur_source") == "altitudes":
                out["hauteur_origine"] += " (altitudes)"
            return


def resolve_row(row, communes_index, bdtopo, anfr, preference="ouvrage"):
    """Pick a structure for one spreadsheet row, and explain the choice."""
    out = dict(row)
    commune = resolve_commune(row["commune"], communes_index)
    if commune is None:
        out["statut"] = "commune not found"
        return out

    out["commune_officielle"] = commune["nom"]
    out["insee"] = commune["code"]
    geom = commune["geom"]
    kind = kind_of(row.get("nature"))

    geo = geocode(row.get("adresse"), commune["nom"], commune["code"])
    if geo:
        gx, gy, out["geocodage"], out["geocodage_score"] = geo
    else:
        cx, cy = to_l93(*commune["centre"]["coordinates"])
        gx, gy, out["geocodage"], out["geocodage_score"] = cx, cy, "(centre commune)", 0.0

    def in_commune(o):
        return geom.contains(Point(o["lon"], o["lat"]))

    # Pools are tried in order, moving down only when the previous one is
    # empty. A water tower in the next commune beats a ground-level reservoir
    # in the right one.
    reservoirs = bdtopo.get("reservoir", [])
    pylones = bdtopo.get("pylone", [])
    tours = [o for o in reservoirs if bdtopo_kind(o["nature"]) == "chateau_eau"]
    sols = [o for o in reservoirs if bdtopo_kind(o["nature"]) == "reservoir_sol"]

    if kind == "pylone":
        pools = [("pylône commune", [o for o in pylones if in_commune(o)]),
                 ("pylône proche", pylones)]
    elif kind == "chateau_eau":
        pools = [("château d'eau commune", [o for o in tours if in_commune(o)]),
                 ("château d'eau proche", tours),
                 ("réservoir au sol commune", [o for o in sols if in_commune(o)])]
    else:
        pools = [("ouvrage commune", [o for o in tours + pylones if in_commune(o)]),
                 ("ouvrage proche", tours + pylones)]

    want = ANFR_CHATEAU_EAU if kind == "chateau_eau" else ANFR_PYLONE

    def attach_anfr(o):
        """ANFR support describing the same structure. ANFR coordinates are in
        whole degrees/minutes/seconds (~30 m), so the tolerance is wide."""
        near = [s for s in anfr
                if dist(o["x"], o["y"], s) < 150 and s["nat_id"] in want]
        o = dict(o)
        if near:
            near.sort(key=lambda s: dist(o["x"], o["y"], s))
            o["anfr"] = near[0]
            o["adr_ratio"] = adresse_ratio(row.get("adresse"), near[0]["adresse"])
        else:
            o["anfr"], o["adr_ratio"] = None, 0.0
        return o

    cands, pool_name = [], None
    for name, pool in pools:
        near = [attach_anfr(o) for o in pool if dist(gx, gy, o) <= RAYON_RECHERCHE_M]
        if near:
            # A matching ANFR address beats proximity to the geocoded point.
            near.sort(key=lambda o: (-(o["adr_ratio"] >= SEUIL_ADRESSE),
                                     -o["adr_ratio"] if o["adr_ratio"] >= SEUIL_ADRESSE else 0,
                                     dist(gx, gy, o)))
            cands, pool_name = near, name
            break

    if not cands:
        out["statut"] = f"no BD TOPO structure ({kind}) within {RAYON_RECHERCHE_M // 1000} km"
    else:
        best = cands[0]
        h, h_src = hauteur_of(best)
        out.update(
            latitude=round(best["lat"], 6), longitude=round(best["lon"], 6),
            hauteur_bdtopo=h, hauteur_source=h_src, z_sol=best["z_sol"],
            nature_bdtopo=best["nature"], cleabs=best["cleabs"],
            dist_geocodage_m=round(dist(gx, gy, best)),
            nb_candidats=len(cands), vivier=pool_name,
        )
        if len(cands) > 1:
            out["ecart_2e_candidat_m"] = round(dist(gx, gy, cands[1]))
        out["statut"] = "ok"
        if best["anfr"]:
            out["hauteur_anfr"] = best["anfr"]["hauteur"]
            out["anfr_dist_m"] = round(dist(best["x"], best["y"], best["anfr"]))
            out["anfr_sup_id"] = best["anfr"]["sup_id"]
            out["anfr_adresse"] = best["anfr"]["adresse"]
            out["adresse_ratio"] = round(best["adr_ratio"], 2)
            if best["adr_ratio"] >= SEUIL_ADRESSE:
                out["choisi_par"] = "adresse ANFR"
        out.setdefault("choisi_par", "proximité géocodage")

    # Without a BD TOPO structure, an ANFR support in the commune is still a lead.
    if out["statut"] != "ok":
        fallback = [s for s in anfr if s["insee"] == commune["code"] and s["nat_id"] in want]
        if fallback:
            fallback.sort(key=lambda s: (-adresse_ratio(row.get("adresse"), s["adresse"]),
                                         dist(gx, gy, s)))
            s = fallback[0]
            out["anfr_adresse"] = s["adresse"]
            out["adresse_ratio"] = round(adresse_ratio(row.get("adresse"), s["adresse"]), 2)
            out.update(
                latitude=round(s["lat"], 6), longitude=round(s["lon"], 6),
                hauteur_anfr=s["hauteur"], anfr_sup_id=s["sup_id"],
                dist_geocodage_m=round(dist(gx, gy, s)), nb_candidats=len(fallback),
                vivier="ANFR seul", statut="ok (ANFR only)",
            )

    if out.get("latitude") is not None:
        hx, hy = to_l93(out["longitude"], out["latitude"])
        out["hauteur_lidar"] = hauteur_lidar(hx, hy)

    choisir_hauteur(out, preference)

    out["alertes"] = alertes_de(out, row)
    return out


def alertes_de(out: dict, row: dict | None = None) -> str:
    """What deserves a human look before launching an hour of computation."""
    row = row if row is not None else out
    alertes = []
    if out.get("statut", "").startswith("ok"):
        if out.get("hauteur_retenue") is None:
            alertes.append("unknown height")
        if out.get("choisi_par") != "adresse ANFR" and (out.get("dist_geocodage_m") or 0) > 800:
            alertes.append(f"structure {out['dist_geocodage_m']} m from the geocoded point")
        e2 = out.get("ecart_2e_candidat_m")
        if (out.get("choisi_par") != "adresse ANFR" and e2 is not None
                and e2 - (out.get("dist_geocodage_m") or 0) < ECART_AMBIGU_M):
            alertes.append("2nd candidate just as close")
        if not row.get("adresse") and out.get("choisi_par") != "adresse ANFR":
            alertes.append("no address in the spreadsheet")
        if out.get("vivier", "").startswith("réservoir au sol"):
            alertes.append("ground-level reservoir, not a tower")
        hr = out.get("hauteur_retenue")
        autres = {"ANFR": out.get("hauteur_anfr"),
                  "BD TOPO": out.get("hauteur_bdtopo"),
                  "lidar": out.get("hauteur_lidar")}
        autres.pop(out.get("hauteur_origine", "").split(" (")[0], None)
        ecarts = [f"{k} {float(v):.0f} m" for k, v in autres.items()
                  if v and hr and abs(float(v) - hr) > 8]
        if ecarts:
            alertes.append(f"{hr} m kept but {', '.join(ecarts)}")
        if out.get("hauteur_lidar") is None and hr:
            alertes.append("height not confirmed by the lidar")
    return " ; ".join(alertes)


# --------------------------------------------------------------------------

def appliquer_installes(resolved: list[dict], path: Path) -> int:
    """Overwrite resolved coordinates with those of cameras already installed.

    The operations JSON is authoritative on position: on a test site, an SDIS
    mast missing from every reference dataset, automatic resolution was 700 m
    off. These sites also form the scenarios' baseline, hence `EN PLACE` in
    the Status column, which `prioritise_sites.py` recognises.

    The JSON's `elevation` field is deliberately ignored: read as an NGF
    altitude it gives −23 m and −48 m on two sites, a camera underground; read
    as a height above ground it would need a 100 m mast on an 82 m tower. It
    does not measure what we need — given four values grouped at 100-102, most
    likely a PTZ tilt angle.

    A site may carry `hauteur_sol_m`, in metres above ground: that value is
    authoritative and replaces any height inferred from the reference data.
    """
    if not path.exists():
        return 0
    installes = json.loads(path.read_text(encoding="utf-8"))
    n = 0
    for site in installes:
        cle = norm(site["name"])
        # A commune may have two rows (water tower and pylon): break the tie on
        # distance to the operations coordinates.
        candidats = [r for r in resolved if norm(r["commune"]) in cle]
        if not candidats:
            print(f"  ! « {site['name']} » matches no row of the spreadsheet",
                  file=sys.stderr)
            continue
        if len(candidats) > 1:
            sx, sy = to_l93(site["lon"], site["lat"])
            def ecart(r):
                if r.get("latitude") is None:
                    return float("inf")
                rx, ry = to_l93(r["longitude"], r["latitude"])
                return math.hypot(sx - rx, sy - ry)
            candidats.sort(key=ecart)
        r = candidats[0]
        if r.get("latitude") is not None:
            ax, ay = to_l93(r["longitude"], r["latitude"])
            bx, by = to_l93(site["lon"], site["lat"])
            r["ecart_resolution_m"] = round(math.hypot(ax - bx, ay - by))
        r.update(latitude=site["lat"], longitude=site["lon"],
                 en_place=True, source_position=f"exploitation ({site['name']})")

        # Beyond this gap the matched structure is not the one carrying the
        # camera: its height is worthless and must be dropped, or the viewshed
        # is computed at the right position with the wrong height.
        if (r.get("ecart_resolution_m") or 0) > 100:
            for k in ("hauteur_bdtopo", "hauteur_anfr", "hauteur_source",
                      "nature_bdtopo", "cleabs", "anfr_sup_id", "z_sol"):
                r.pop(k, None)
            r["statut"] = "operations position, structure not identified"
        # The lidar is re-measured at the real position in every case.
        hx, hy = to_l93(site["lon"], site["lat"])
        r["hauteur_lidar"] = hauteur_lidar(hx, hy)

        if site.get("hauteur_sol_m") is not None:
            r["hauteur_exploitation"] = float(site["hauteur_sol_m"])
            if site.get("hauteur_note"):
                r["hauteur_note"] = site["hauteur_note"]
        n += 1
    return n


REVIEW_COLS = [
    "ligne", "num", "etape", "commune", "commune_officielle", "insee", "nature",
    "adresse", "zone", "gestionnaire", "avancement", "statut", "alertes",
    "latitude", "longitude", "hauteur_retenue", "hauteur_origine",
    "hauteur_exploitation", "hauteur_anfr", "hauteur_bdtopo", "hauteur_lidar",
    "hauteur_source", "hauteur_note",
    "z_sol", "nature_bdtopo",
    "en_place", "source_position", "ecart_resolution_m", "vivier", "choisi_par", "anfr_adresse", "adresse_ratio", "dist_geocodage_m", "ecart_2e_candidat_m", "nb_candidats",
    "geocodage", "geocodage_score", "anfr_dist_m", "anfr_sup_id", "cleabs",
]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, type=Path,
                    help="the partner's spreadsheet of high points")
    ap.add_argument("--region", required=True, help="folder under data/")
    ap.add_argument("--dept", default=None,
                    help="département code, e.g. 77 (default: \"dept\" in "
                         "data/<region>/scenarios.json)")
    ap.add_argument("--anfr", type=Path,
                    default=Path.home() / "Downloads" / "SUP_SUPPORT.geojson",
                    help="ANFR SUP_SUPPORT.geojson export (default: ~/Downloads)")
    ap.add_argument("--etapes", default="vert,orange,jaune",
                    help="stages (row colours) kept in sites.csv "
                         "(default: all but blanc)")
    ap.add_argument("--installes", type=Path, default=None,
                    help="JSON of cameras already installed (default: "
                         "data/<region>/sites_installes.json if present)")
    ap.add_argument("--hauteur", default="ouvrage",
                    choices=("ouvrage", "anfr", "lidar"),
                    help="preferred height source (default: ouvrage, "
                         "i.e. BD TOPO, then lidar, then ANFR)")
    args = ap.parse_args()

    args.dept = region_dept(args.region, args.dept)
    if not args.dept:
        raise SystemExit(f"❌ no département: set \"dept\" in "
                         f"data/{args.region}/scenarios.json, or pass --dept")
    outdir = REPO / "data" / args.region
    outdir.mkdir(parents=True, exist_ok=True)
    COMMUNE_ALIAS.update(load_commune_alias(outdir / "commune_alias.json"))
    if COMMUNE_ALIAS:
        print(f"  {len(COMMUNE_ALIAS)} commune alias(es) from commune_alias.json")

    print(f"▶ {args.xlsx.name}")
    rows = read_xlsx(args.xlsx)
    print(f"  {len(rows)} high-point rows")

    print("▶ reference datasets")
    communes_index, communes = load_communes(args.dept)
    bdtopo = load_bdtopo(communes)
    anfr = load_anfr(args.anfr, args.dept) if args.anfr.exists() else []
    if not anfr:
        print(f"  ! ANFR export missing ({args.anfr}) — no height cross-check",
              file=sys.stderr)

    print("▶ resolution")
    resolved = [resolve_row(row, communes_index, bdtopo, anfr, args.hauteur)
                for row in rows]

    installes = args.installes or (outdir / "sites_installes.json")
    n_installes = appliquer_installes(resolved, installes)
    for r in resolved:
        if r.get("en_place"):
            # height and alerts recomputed on what is left after repositioning
            for k in ("hauteur_retenue", "hauteur_origine"):
                r.pop(k, None)
            choisir_hauteur(r, args.hauteur)
            r["alertes"] = alertes_de(r)
            if r.get("hauteur_retenue") is None:
                r["alertes"] = "; ".join(filter(None, [
                    "unknown height, ask the partner", r["alertes"]]))
    # Printed only now: installed sites had their position, and sometimes their
    # height, replaced above, and showing the intermediate value would mislead.
    for r in resolved:
        h = r.get("hauteur_retenue")
        flag = "≡" if r.get("en_place") else " "
        print(f"  {flag}[{r['etape']:6s}] {str(r['commune'])[:22]:22s} "
              f"{r['statut'][:30]:30s} h={str(h) + ' m' if h else '?':8s} "
              f"{r['alertes']}")
    if n_installes:
        print(f"\n  ≡ {n_installes} site(s) repositioned from "
              f"{installes.name} and marked EN PLACE")

    review = outdir / "resolution.csv"
    with review.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=REVIEW_COLS, delimiter=";",
                           extrasaction="ignore")
        w.writeheader()
        w.writerows(resolved)

    keep = {e.strip() for e in args.etapes.split(",")}
    sites = outdir / "sites.csv"
    n = 0
    with sites.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Name", "Latitude", "Longitude", "Owner", "Height", "Status"])
        for r in resolved:
            if r["etape"] not in keep or not r.get("latitude"):
                continue
            h = r.get("hauteur_retenue")
            if h is None:
                continue
            name = f"{r.get('num') or ''} {r['commune']}".strip().replace(";", ",")
            # Owner keeps the manager — it carries the access cost, which
            # separates the scenarios. Whether a camera is already installed is
            # another piece of information: it goes in Status, otherwise
            # « EN PLACE » in Owner would erase the delegate or the operator.
            owner = str(r.get("gestionnaire") or "").replace(";", ",")
            w.writerow([name, r["latitude"], r["longitude"], owner,
                        round(float(h), 1),
                        "EN PLACE" if r.get("en_place") else ""])
            n += 1

    flagged = sum(1 for r in resolved if r.get("alertes"))
    print(f"\n✓ {review.relative_to(REPO)}  ({len(resolved)} rows, "
          f"{flagged} to check — column « alertes »)")
    print(f"✓ {sites.relative_to(REPO)}  ({n} sites, stages: {', '.join(sorted(keep))})")
    print(f"\n  Heights: preference « {args.hauteur} ». This is the height of the")
    print("  STRUCTURE, not of the camera. Check with audit_sites.py and adjust")
    print("  before running the viewsheds.")


if __name__ == "__main__":
    main()
